"""
简历解析服务
支持 PDF（pypdf）和 Word（python-docx）格式
"""
import os
from pypdf import PdfReader
from docx import Document
from ..config import ALLOWED_EXTENSIONS


def parse_resume(file_path: str, filename: str) -> str:
    """
    根据文件类型解析简历文本
    Args:
        file_path: 文件完整路径
        filename: 文件名（用于判断扩展名）
    Returns:
        解析后的文本内容
    """
    ext = filename.rsplit(".", 1)[-1].lower()

    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"不支持的文件格式：.{ext}，仅支持 {ALLOWED_EXTENSIONS}")

    if ext == "pdf":
        return _parse_pdf(file_path)
    elif ext == "docx":
        return _parse_docx(file_path)
    elif ext == "txt":
        return _parse_txt(file_path)


def _parse_pdf(file_path: str) -> str:
    """使用 pypdf 解析 PDF 文件"""
    text_parts = []
    try:
        reader = PdfReader(file_path)
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
        return "\n".join(text_parts).strip()
    except Exception as e:
        raise RuntimeError(f"PDF 解析失败：{str(e)}")


def _parse_docx(file_path: str) -> str:
    """使用 python-docx 解析 Word 文件（段落 + 表格）"""
    text_parts = []
    try:
        doc = Document(file_path)
        # 1. 普通段落
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text.strip())
        # 2. 表格单元格（表格式简历往往只有 tables，没有段落）
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    text_parts.append("  ".join(row_cells))
        return "\n".join(text_parts).strip()
    except Exception as e:
        raise RuntimeError(f"Word 文件解析失败：{str(e)}")



def _parse_txt(file_path: str) -> str:
    """解析纯文本文件"""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read().strip()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="gbk") as f:
            return f.read().strip()
    except Exception as e:
        raise RuntimeError(f"文本文件解析失败：{str(e)}")
